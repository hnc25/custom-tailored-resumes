#!/usr/bin/env python3
"""
Regression tests for the tailoring workflow, on a SYNTHETIC candidate (tests/fixtures).
No real candidate data is read or written. Run from the repo root:

    python -m unittest discover -s tests -v

What these tests can and cannot show: they prove the recorded rules behave as designed on
known cases. They do not prove that a real deliverable is factually correct. The last test
in each group that is named "limit_..." pins down a known blind spot on purpose.
"""

import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import docx

REPO = Path(__file__).resolve().parent.parent
FIX = Path(__file__).resolve().parent / "fixtures"
sys.path.insert(0, str(REPO / "scripts"))

import ats_check  # noqa: E402
import evidence_map  # noqa: E402
import qa_check  # noqa: E402

CORE_LINES = [
    "JORDAN AVERY, MBA",
    "NORTHWIND LOGISTICS — Senior Product Manager",
    "Led the shipment-tracking portal for 40+ enterprise shippers.",
    "LUMEN RETAIL GROUP — Product Manager",
    "Owned the loyalty app roadmap, with $4M in spend.",
]


def make_root(tmp, rules=None):
    """A complete candidate-owned file set for the synthetic candidate, under tmp."""
    ref = Path(tmp) / "reference"
    ref.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    for line in CORE_LINES:
        d.add_paragraph(line)
    d.save(ref / "core-resume.docx")
    shutil.copy(FIX / "background-notes.md", ref / "background-notes.md")
    rules = rules or json.loads((FIX / "evidence_rules.json").read_text(encoding="utf-8"))
    (ref / "evidence_rules.json").write_text(json.dumps(rules), encoding="utf-8")
    return Path(tmp)


def rules_of(errors):
    return [e[0] for e in errors]


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.rules = json.loads((FIX / "evidence_rules.json").read_text(encoding="utf-8"))
        self.gate = qa_check.Gate(root=make_root(self._tmp.name))
        self.emap = json.loads((FIX / "evidence_map_AcmeFreight.json").read_text(encoding="utf-8"))
        self.app = json.loads((FIX / "application_AcmeFreight.json").read_text(encoding="utf-8"))

    def gate_with(self, rules):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        return qa_check.Gate(root=make_root(tmp.name, rules))

    def check(self, lines, emap=None, letter=False, gate=None, notices=None):
        fn = qa_check.check_cover_letter if letter else qa_check.check_paragraphs
        return fn(lines, gate=gate or self.gate, evidence_map=emap if emap is not None else {}, notices=notices)


class SupportedParaphrase(Base):
    """1. A professional paraphrase of a recorded activity passes without the exact source keyword."""

    LINE = ["NORTHWIND LOGISTICS", "Ran weekly funnel analysis in SQL of carrier onboarding."]

    def test_keyword_absent_from_sources(self):
        self.assertNotIn("funnel", self.gate.source_text().lower())

    def test_undocumented_term_without_review_needs_review_not_gap(self):
        errs = self.check(self.LINE)
        self.assertEqual(rules_of(errs), ["EVIDENCE REVIEW REQUIRED"])
        self.assertIn("not a confirmed gap", errs[0][2])

    def test_paraphrase_grounded_in_recorded_activity_passes_with_notice(self):
        notices = []
        self.assertEqual(self.check(self.LINE, self.emap, notices=notices), [])
        self.assertEqual(len(notices), 1)
        self.assertIn("not checked automatically", notices[0][2])

    def test_label_without_activity_is_rejected(self):
        emap = copy.deepcopy(self.emap)
        emap["rows"][0]["activity"] = ""
        self.assertEqual(rules_of(self.check(self.LINE, emap)), ["EVIDENCE REVIEW REQUIRED"])
        self.assertTrue(any("does not replace the activity" in e for e in evidence_map.validate_map(emap, gate=self.gate)))

    def test_invented_source_quote_is_rejected(self):
        emap = copy.deepcopy(self.emap)
        emap["rows"][0]["source_quote"] = "I ran funnel analysis every week for three years."
        self.assertEqual(rules_of(self.check(self.LINE, emap)), ["EVIDENCE REVIEW REQUIRED"])
        self.assertTrue(any("not in the candidate's sources" in e for e in evidence_map.validate_map(emap, gate=self.gate)))

    def test_review_of_one_term_does_not_cover_another(self):
        errs = self.check(["NORTHWIND LOGISTICS", "Ran funnel analysis and A/B testing on onboarding."], self.emap)
        self.assertEqual(rules_of(errs), ["EVIDENCE REVIEW REQUIRED", "UNANSWERED CLAIM"])
        self.assertIn("A/B testing", errs[0][2])


class QuestionVersusGap(Base):
    """2. An unanswered question is not a confirmed gap, and neither can carry a claim."""

    def test_fixture_map_is_complete(self):
        self.assertEqual(evidence_map.validate_map(self.emap, self.app["keywords"], self.gate), [])

    def test_missing_documentation_cannot_be_called_a_gap(self):
        emap = copy.deepcopy(self.emap)
        row = emap["rows"][2]  # A/B testing: nothing in the sources either way
        row.update(classification="confirmed_gap", gap_label="A/B testing")
        row.pop("question")
        errs = evidence_map.validate_map(emap, gate=self.gate)
        self.assertTrue(any("open_question" in e and "undocumented" in e for e in errs), errs)

    def test_gap_with_no_confirmation_at_all_is_rejected(self):
        emap = copy.deepcopy(self.emap)
        emap["rows"][3].pop("gap_label")
        self.assertTrue(any("nothing shows the candidate confirmed" in e for e in evidence_map.validate_map(emap, gate=self.gate)))

    def test_open_question_cannot_carry_a_claim(self):
        emap = copy.deepcopy(self.emap)
        emap["rows"][2]["claim"] = "Ran A/B tests on onboarding"
        self.assertTrue(any("cannot carry a claim" in e for e in evidence_map.validate_map(emap, gate=self.gate)))

    def test_questions_are_collected_together_and_stop_the_build(self):
        self.assertEqual([q[0] for q in evidence_map.open_questions(self.emap)], ["Run A/B tests and controlled experiments"])
        p = Path(self._tmp.name) / "evidence_map_AcmeFreight.json"
        p.write_text(json.dumps(self.emap), encoding="utf-8")
        with self.assertRaises(SystemExit) as cm:
            evidence_map.require(p, self.app["keywords"], self.gate)
        self.assertIn("Open questions for the candidate", str(cm.exception))
        self.assertTrue(evidence_map.require(p, self.app["keywords"], self.gate, questions_asked=True))

    def test_asking_a_question_does_not_authorise_its_claim(self):
        """The term has NO entry in the rules file, so only the evidence map stands between an
        unanswered question and a claim. Marking questions as asked must not open that door."""
        emap = copy.deepcopy(self.emap)
        emap["rows"].append({"requirement": "Own pricing strategy", "importance": "must", "jd_meaning": "Set and test prices",
                             "classification": "open_question", "question": "Have you owned pricing decisions?",
                             "keywords": ["pricing strategy"]})
        p = Path(self._tmp.name) / "evidence_map_AcmeFreight.json"
        p.write_text(json.dumps(emap), encoding="utf-8")
        asked = evidence_map.require(p, gate=self.gate, questions_asked=True)
        claim = ["NORTHWIND LOGISTICS", "Owned pricing strategy for the carrier portal."]
        self.assertEqual(self.check(claim), [])  # blind without the map: nothing in the rules names it
        errs = self.check(claim, asked)
        self.assertEqual(rules_of(errs), ["UNANSWERED CLAIM"])
        self.assertIn("does not authorise", errs[0][2])
        letter = ["Dear Hiring Manager,", "I owned pricing strategy for the carrier portal."]
        self.assertIn("UNANSWERED CLAIM", rules_of(self.check(letter, asked, letter=True)))

    def test_supported_content_proceeds_while_a_question_is_open(self):
        lines = ["NORTHWIND LOGISTICS", "Ran weekly funnel analysis in SQL of carrier onboarding.", "Scheduled weekly carrier reports in Airflow."]
        self.assertEqual(self.check(lines, self.emap), [])

    def test_confirmed_gap_keyword_is_blocked_but_keyword_only_sense_is_allowed(self):
        self.assertIn("UNANSWERED CLAIM", rules_of(self.check(["LUMEN RETAIL GROUP", "Handled sales management reporting."], self.emap)))
        self.assertEqual(self.check(["LUMEN RETAIL GROUP", "Presented quarterly pipeline reviews of the product roadmap."], self.emap), [])

    def test_confirmed_gap_blocks_even_with_a_review_row(self):
        emap = copy.deepcopy(self.emap)
        emap["rows"][0]["reviewed_terms"].append("sales team")
        errs = self.check(["NORTHWIND LOGISTICS", "Managed the regional sales team."], emap)
        self.assertEqual(rules_of(errs), ["KNOWN-GAP CONTRADICTION"])

    def test_unmapped_must_have_keyword_is_reported(self):
        kw = {"must_have": self.app["keywords"]["must_have"] + ["pricing strategy"]}
        self.assertTrue(any("pricing strategy" in e for e in evidence_map.validate_map(self.emap, kw, self.gate)))


class ConfirmationResolvesRestriction(Base):
    """3. A new confirmation lifts the restriction it resolves, and only that one."""

    def stale_rules(self):
        rules = copy.deepcopy(self.rules)
        rules["gaps"][3].update(label="Kafka / Airflow", claim=r"\b(Kafka|Airflow)\b")
        return rules

    def test_stale_restriction_is_detected(self):
        errs = evidence_map.check_store(self.gate_with(self.stale_rules()))
        self.assertTrue(any("stale restriction" in e and "Airflow" in e for e in errs), errs)

    def test_resolved_store_is_consistent(self):
        self.assertEqual(evidence_map.check_store(self.gate), [])

    def test_confirmed_term_is_available_without_any_evidence_map(self):
        self.assertEqual(self.check(["NORTHWIND LOGISTICS", "Scheduled weekly carrier reports in Airflow."]), [])

    def test_unresolved_half_stays_restricted(self):
        errs = self.check(["NORTHWIND LOGISTICS", "Scheduled reports in Airflow and streamed events through Kafka."])
        self.assertEqual(rules_of(errs), ["EVIDENCE REVIEW REQUIRED"])
        self.assertIn("Kafka", errs[0][2])

    def test_confirmation_does_not_spread_to_another_employer(self):
        errs = self.check(["LUMEN RETAIL GROUP", "Scheduled loyalty reports in Airflow."])
        self.assertEqual(rules_of(errs), ["EMPLOYER SCOPE"])

    def test_confirmation_without_saved_answer_is_rejected(self):
        rules = copy.deepcopy(self.rules)
        rules["confirmed_terms"][0]["source_quote"] = "I used Airflow everywhere."
        self.assertTrue(any("Save the candidate's full answer" in e for e in evidence_map.check_store(self.gate_with(rules))))

    def test_confirmed_status_needs_source_and_date(self):
        rules = copy.deepcopy(self.rules)
        del rules["gaps"][0]["confirmed_on"]
        self.assertTrue(any("confirmed_on" in e for e in evidence_map.check_store(self.gate_with(rules))))


class MisleadingKeyword(Base):
    """4. A keyword that matches the text in a different sense is not substantive coverage."""

    TEXT = ("Presented quarterly pipeline reviews of the product roadmap. Ran weekly funnel analysis in SQL. "
            "On-time delivery improved. Scheduled reports in Airflow.")

    def fit(self):
        kw = ats_check.check_keywords(self.TEXT, self.app["keywords"])
        return kw, ats_check.role_fit_report(kw, self.emap)

    def test_keyword_score_counts_the_misleading_match(self):
        kw, _ = self.fit()
        self.assertIn("pipeline", [t for t, _ in kw["must_have_matched"]])

    def test_role_fit_does_not(self):
        kw, fit = self.fit()
        self.assertEqual(fit["matched_keyword_only"], ["pipeline"])
        self.assertLess(fit["must_have_substantive_pct"], kw["must_have_pct"])

    def test_role_fit_names_open_questions_and_gaps_separately(self):
        _, fit = self.fit()
        must = fit["requirements"]["must"]
        self.assertEqual(len(must["open_question"]), 1)
        self.assertEqual(len(must["confirmed_gap"]), 1)
        self.assertIn("does not establish", fit["limits"])


class CausalAndOwnership(Base):
    """5. Unsupported causal and ownership claims that match a recorded rule are blocked."""

    def test_analysis_joined_to_result_figure_is_blocked(self):
        errs = self.check(["NORTHWIND LOGISTICS", "Ran weekly funnel analysis that improved on-time delivery from 82% to 91%."], self.emap)
        self.assertIn("UNSUPPORTED LINK", rules_of(errs))

    def test_same_facts_in_separate_sentences_pass(self):
        lines = ["NORTHWIND LOGISTICS", "Ran weekly funnel analysis in SQL of carrier onboarding.",
                 "Wrote the requirements for a routing release, after which on-time delivery improved from 82% to 91%."]
        self.assertEqual(self.check(lines, self.emap), [])

    def test_ownership_the_candidate_denied_is_blocked(self):
        errs = self.check(["NORTHWIND LOGISTICS", "Built the routing models behind the release."])
        self.assertEqual(rules_of(errs), ["CONFIRMED BOUNDARY"])

    def test_unsourced_figure_is_blocked(self):
        self.assertEqual(rules_of(self.check(["Cut onboarding time by 37%."])), ["UNSOURCED FIGURE"])

    def test_limit_unpatterned_causal_claim_is_not_caught(self):
        """Known blind spot, stated in every report: a causal claim worded outside the recorded
        patterns passes the gate. Only a reader comparing it with the evidence map catches it."""
        line = ["NORTHWIND LOGISTICS", "Weekly onboarding reviews were the reason delivery reliability rose."]
        self.assertEqual(self.check(line), [])
        for text in (qa_check.LIMITS, evidence_map.LIMITS):
            self.assertRegex(text, r"do(es)? not")


class ShippedFilesAreNeutral(unittest.TestCase):
    """6a. The reusable scripts carry no candidate: everything comes from candidate-owned files."""

    def test_scripts_run_against_a_bare_candidate_root(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = make_root(tmp)
            (root / "reference" / "evidence_rules.json").unlink()
            gate = qa_check.Gate(root=root)
            self.assertEqual(gate.gaps, [])
            self.assertEqual(rules_of(qa_check.check_paragraphs(["He led the portal for 40+ shippers."], gate=gate, evidence_map={})), ["THIRD-PERSON VOICE"])

    def test_fixtures_use_reserved_example_identity(self):
        app = (FIX / "application_AcmeFreight.json").read_text(encoding="utf-8")
        self.assertIn("@example.com", app)
        self.assertIn("555-01", app)


class EndToEnd(unittest.TestCase):
    """The three deliverables build from one definition, using only candidate-owned files."""

    def test_one_definition_builds_three_consistent_files(self):
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("Pillow not installed")
        if not (REPO / "node_modules" / "docx").exists() or not shutil.which("node"):
            self.skipTest("Node renderer not installed (run npm install)")
        with tempfile.TemporaryDirectory() as tmp:
            root = make_root(tmp)
            shutil.copytree(REPO / "scripts", root / "scripts", ignore=shutil.ignore_patterns("__pycache__", "job*"))
            src = (root / "scripts" / "build_professional.py").read_text(encoding="utf-8")
            photo = "HC_headshot.jpg" if "HC_headshot.jpg" in src else "photo.jpg"
            Image.new("RGB", (500, 600), (120, 130, 140)).save(root / "reference" / photo)
            shutil.copy(FIX / "application_AcmeFreight.json", root)
            shutil.copy(FIX / "evidence_map_AcmeFreight.json", root)
            cmd = [sys.executable, str(root / "scripts" / "build_application.py"), str(root / "application_AcmeFreight.json")]
            env = {**os.environ, "NODE_PATH": str(REPO / "node_modules"), "PYTHONIOENCODING": "utf-8"}
            blocked = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", env=env, cwd=root)
            self.assertNotEqual(blocked.returncode, 0)
            self.assertIn("Open questions for the candidate", blocked.stderr)
            # asked is not answered: the same definition plus the unanswered claim must build nothing
            app = json.loads((root / "application_AcmeFreight.json").read_text(encoding="utf-8"))
            app["experience"][0]["bullets"].append("Ran A/B testing on carrier onboarding emails.")
            (root / "application_AcmeFreight.json").write_text(json.dumps(app), encoding="utf-8")
            unanswered = subprocess.run(cmd + ["--questions-asked"], capture_output=True, text=True, encoding="utf-8", env=env, cwd=root)
            self.assertNotEqual(unanswered.returncode, 0)
            self.assertIn("UNANSWERED CLAIM", unanswered.stderr)
            self.assertFalse((root / "output").exists() and any((root / "output").iterdir()))
            shutil.copy(FIX / "application_AcmeFreight.json", root)
            done = subprocess.run(cmd + ["--questions-asked"], capture_output=True, text=True, encoding="utf-8", env=env, cwd=root)
            self.assertEqual(done.returncode, 0, done.stderr[-3000:])
            self.assertIn("Role Fit", done.stdout)
            self.assertIn("Limits of this gate", done.stdout)
            out = root / "output"
            names = ["Resume_AcmeFreight_2025-03-12.docx", "Professional_AcmeFreight_2025-03-12.docx", "Professional_CoverLetter_AcmeFreight_2025-03-12.docx"]
            texts = [" ".join(qa_check._doc_text(docx.Document(str(out / n)))) for n in names]
            bullet = "Ran weekly funnel analysis in SQL of carrier onboarding"
            self.assertIn(bullet, texts[0])
            self.assertIn(bullet, texts[1])
            for t in texts:
                self.assertIn("Jordan Avery".lower(), t.lower())


if __name__ == "__main__":
    unittest.main(verbosity=2)
