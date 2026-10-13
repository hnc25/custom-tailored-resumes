#!/usr/bin/env python3
"""
build_docx.py — Build ATS-safe resume and cover letter .docx files in a plain
single-column style: single-size 9.5pt Calibri, bold for hierarchy
(name/section headers/company names/job titles), List Bullet style for bullets,
US Letter page, narrow margins, no tables/textboxes/images/headers/footers.

Usage: import as a module from a per-job build script.
"""

import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from qa_check import enforce

FONT_NAME = "Calibri"
FONT_SIZE = Pt(9.5)


def _set_run(run, bold=False, italic=False):
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    run.bold = bold
    run.italic = italic
    return run


def _new_doc():
    d = docx.Document()
    section = d.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.3125)
    section.bottom_margin = Inches(0.4375)
    # Ensure header/footer stay empty
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    normal = d.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = FONT_SIZE
    return d


def add_para(d, runs, bullet=False, space_after=0):
    """runs: list of (text, bold) tuples, or a single string (treated as one non-bold run)."""
    if isinstance(runs, str):
        runs = [(runs, False)]
    style = "List Bullet" if bullet else None
    p = d.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for text, bold in runs:
        r = p.add_run(text)
        _set_run(r, bold=bold)
    return p


def add_blank(d):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("")
    _set_run(r)
    return p


def build_resume(path, header_lines, blocks):
    """
    header_lines: list of run-lists for the top identity block (name, contact, title, tagline)
    blocks: list of ("heading"|"text"|"bullet"|"blank", content) tuples where content
            for heading/text/bullet is a runs list (see add_para).
    """
    d = _new_doc()
    for line in header_lines:
        add_para(d, line)
    add_blank(d)
    for kind, content in blocks:
        if kind == "blank":
            add_blank(d)
        elif kind == "heading":
            add_para(d, [(content, True)] if isinstance(content, str) else content)
        elif kind == "text":
            add_para(d, content)
        elif kind == "bullet":
            add_para(d, content, bullet=True)
    enforce(d, path)  # mandatory QA gate — exits without saving on any error
    d.save(path)


def build_cover_letter(path, header_lines, date_line, recipient_lines, salutation, paragraphs, closing_lines):
    d = _new_doc()
    for line in header_lines:
        add_para(d, line)
    add_blank(d)
    add_para(d, date_line)
    add_blank(d)
    for line in recipient_lines:
        add_para(d, line)
    add_blank(d)
    add_para(d, salutation)
    add_blank(d)
    for para in paragraphs:
        add_para(d, para)
        add_blank(d)
    for line in closing_lines:
        add_para(d, line)
    enforce(d, path)  # mandatory QA gate — exits without saving on any error
    d.save(path)
